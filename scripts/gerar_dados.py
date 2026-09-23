# -*- coding: utf-8 -*-
"""Extrai os dados de Brasil_2016_2019_2022_2025.xlsx para docs/data/dados.js.

Uso:  python scripts/gerar_dados.py
Gera um arquivo JS (window.DADOS = {...}) para que o site funcione tanto no
GitHub Pages / Render quanto abrindo o index.html direto do disco.
"""
import json
import pathlib
import re

import openpyxl

RAIZ = pathlib.Path(__file__).resolve().parents[1]
XLSX = RAIZ / "Brasil_2016_2019_2022_2025.xlsx"
SAIDA = RAIZ / "docs" / "data" / "dados.js"

ANOS = [2016, 2019, 2022, 2025]

# Abas temáticas: chave usada no site -> nome da aba
ABAS = {
    "pib": "PIB setorial",
    "receitas": "Receitas públicas",
    "despesas": "Despesas públicas",
    "divida": "Dívida e juros",
    "loa": "LOA e execução orçamentária",
    "social": "Saúde, educação e previdência",
    "seguranca": "Segurança e polít. afirmativas",
    "afirmativas": "Políticas afirmativas e apoio",
    "fomento": "BNDES, Plano Safra e P&D",
    "emprego": "Emprego",
    "externo": "Comércio exterior",
    "privado": "Invest. privados e industriais",
    "infra": "Infraestrutura federal",
}

# Correções de erros da planilha: no script de origem as chaves "imp" e "adm"
# colidiram (importações × impostos; admissões × administração pública), e os
# valores de 2016/2022 dessas duas linhas do PIB setorial ficaram trocados.
# Valores corretos (IBGE SIDRA 1846) vêm do próprio build_2016_2019_2022_2025.py.
CORRECOES = {
    ("pib", "1.2"): {2016: 849.506, 2022: 1343.201},
    ("pib", "4.6"): {2016: 945.120, 2022: 1365.840},
}
FONTE_CORRIGIDA = {
    ("pib", "1.2"): "IBGE/SIDRA 1846 (S01)",
    ("pib", "4.6"): "IBGE/SIDRA 1846 (S01)",
}


def num(v):
    if isinstance(v, (int, float)):
        return round(float(v), 6)
    return None


def txt(v):
    if v is None:
        return ""
    return str(v).strip()


def parametros(wb):
    ws = wb["Resumo executivo"]
    lin = {txt(r[0]): r[1:5] for r in ws.iter_rows(min_row=4, max_row=9, values_only=True) if r[0]}
    def pega(prefixo):
        for k, v in lin.items():
            if k.startswith(prefixo):
                return [num(x) for x in v]
        raise KeyError(prefixo)
    return {
        "anos": ANOS,
        "pib": pega("PIB (R$ bi"),
        "ipca": pega("IPCA"),
        "fator": pega("Fator IPCA"),
        "ptax": pega("PTAX"),
        "pib_usd": pega("PIB em US$"),
    }


def resumo(wb):
    ws = wb["Resumo executivo"]
    rows = list(ws.iter_rows(values_only=True))
    titulo, sub = txt(rows[0][0]), txt(rows[1][0])
    sintese, atual = [], None
    for r in rows:
        c = txt(r[0])
        if c.startswith("■"):
            atual = {"titulo": c.lstrip("■ ").strip(), "texto": ""}
            sintese.append(atual)
        elif atual is not None and c and not atual["texto"]:
            atual["texto"] = c
    legenda = []
    for r in rows:
        if txt(r[0]) in ("OFICIAL", "ESTIMADO", "APROXIMAÇÃO", "IMPRENSA/SETORIAL", "NÃO LOCALIZADO"):
            legenda.append({"classe": txt(r[0]), "descricao": txt(r[1])})
    return {"titulo": titulo, "subtitulo": sub, "sintese": sintese, "legenda": legenda}


def aba_tematica(wb, chave, nome):
    ws = wb[nome]
    rows = list(ws.iter_rows(values_only=True))
    cab = [txt(c) for c in rows[2]]
    idx = {c: i for i, c in enumerate(cab)}
    linhas, notas = [], []
    em_notas = False
    for r in rows[3:]:
        c0 = txt(r[0])
        if c0 == "Notas metodológicas":
            em_notas = True
            continue
        if em_notas:
            if c0:
                notas.append(re.sub(r"^\d+\.\s*", "", c0))
            continue
        if not c0 or r[1] is None:
            continue
        valores = [r[idx[f"Brasil {a}"]] for a in ANOS]
        v = [num(x) for x in valores]
        faltas = [txt(x) if num(x) is None else "" for x in valores]
        fontes = [txt(r[idx[f"Fonte {a}"]]) for a in ANOS]
        corr = CORRECOES.get((chave, c0))
        if corr:
            for i, a in enumerate(ANOS):
                if a in corr:
                    v[i] = corr[a]
                    fontes[i] = FONTE_CORRIGIDA[(chave, c0)]
        linhas.append({
            "cod": c0,
            "nome": txt(r[1]),
            "estagio": txt(r[2]),
            "ambito": txt(r[3]),
            "unidade": txt(r[4]),
            "v": v,
            "falta": faltas,
            "fontes": fontes,
            "classe": txt(r[idx["Classificação do dado"]]) if "Classificação do dado" in idx else "",
            "obs": txt(r[idx["Observações metodológicas"]]) if "Observações metodológicas" in idx else "",
            "corrigido": bool(corr),
        })
    return {"titulo": txt(rows[0][0]), "subtitulo": txt(rows[1][0]), "linhas": linhas, "notas": notas}


def fontes(wb):
    ws = wb["Fontes e metodologia"]
    rows = list(ws.iter_rows(values_only=True))
    lista, tentativas, regras = [], [], []
    modo = "fontes"
    for r in rows[3:]:
        c0 = txt(r[0])
        if c0.startswith("Dados não localizados"):
            modo = "tent"
            continue
        if c0.startswith("Regras metodológicas"):
            modo = "regras"
            continue
        if modo == "fontes" and re.match(r"^S\d+", c0):
            lista.append({
                "id": c0, "titulo": txt(r[1]), "instituicao": txt(r[2]), "url": txt(r[3]),
                "publicacao": txt(r[4]), "referencia": txt(r[5]), "indicador": txt(r[6]),
                "obs": txt(r[8]), "confiabilidade": txt(r[9]),
            })
        elif modo == "tent" and txt(r[1]) and txt(r[1]) != "Item":
            tentativas.append({"item": txt(r[1]), "tentadas": txt(r[3]), "resultado": txt(r[8])})
        elif modo == "regras" and c0:
            regras.append(re.sub(r"^\d+\.\s*", "", c0))
    return {"fontes": lista, "tentativas": tentativas, "regras": regras}


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    dados = {
        "parametros": parametros(wb),
        "resumo": resumo(wb),
        "abas": {k: aba_tematica(wb, k, n) for k, n in ABAS.items()},
        "metodologia": fontes(wb),
    }
    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    corpo = json.dumps(dados, ensure_ascii=False, separators=(",", ":"))
    SAIDA.write_text("// Gerado por scripts/gerar_dados.py — não editar à mão\nwindow.DADOS = " + corpo + ";\n", encoding="utf-8")
    n = sum(len(a["linhas"]) for a in dados["abas"].values())
    print(f"{SAIDA.relative_to(RAIZ)}: {n} linhas, {len(dados['metodologia']['fontes'])} fontes")


if __name__ == "__main__":
    main()
